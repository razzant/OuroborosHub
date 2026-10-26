"""office_reader through plugin.register with a fake PluginAPI; every fixture is generated here.

The fixture generators (python-docx, python-pptx, openpyxl, xlwt) are declared test dependencies of the
CI suite, so no format is skipped: the real converters run for DOCX, PPTX, XLSX, XLS and a hand-built MSG.
"""
from __future__ import annotations

import concurrent.futures
import importlib.util
import json
import os
import pathlib
import struct
import sys
import tempfile
import zipfile

import docx
import openpyxl
import pptx
import pytest
import xlwt

SKILL_DIR = pathlib.Path(__file__).resolve().parents[1]


class FakeAPI:
    def __init__(self, state_dir: pathlib.Path | None):
        self.state_dir = state_dir
        self.tools = {}
        self.unload = None

    def register_tool(self, name, handler, *, description, schema, timeout_sec=60):
        assert len(name) <= 24 and name.replace("_", "").isalnum()
        assert 0 < timeout_sec <= 300 and description and schema["type"] == "object"
        self.tools[name] = handler

    def on_unload(self, callback):
        self.unload = callback

    def get_state_dir(self):
        if self.state_dir is None:
            raise RuntimeError("no state directory for this host")
        return str(self.state_dir)


@pytest.fixture
def env(tmp_path, monkeypatch):
    """Isolated roots: home, temp dir and data root all under tmp_path; `outside` is none of them."""
    data = tmp_path / "data"
    state = data / "state" / "skills" / "office_reader"
    home, temp, outside = tmp_path / "home", tmp_path / "tmpdir", tmp_path / "outside"
    for directory in (data / "uploads", state, home / "Documents", temp, outside):
        directory.mkdir(parents=True)
    monkeypatch.setenv("HOME", str(home))
    monkeypatch.delenv("OUROBOROS_DATA_DIR", raising=False)
    monkeypatch.delenv("OUROBOROS_SKILL_STATE_DIR", raising=False)
    spec = importlib.util.spec_from_file_location("office_reader_under_test", SKILL_DIR / "plugin.py")
    plugin = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = plugin
    spec.loader.exec_module(plugin)
    monkeypatch.setattr(tempfile, "tempdir", str(temp))  # honoured by the plugin's read roots and by openpyxl
    api = FakeAPI(state)
    plugin.register(api)
    yield {"plugin": plugin, "api": api, "uploads": data / "uploads", "home": home, "temp": temp,
           "outside": outside, "state": state}
    api.unload()
    sys.modules.pop(spec.name, None)


class FakeCtx:
    """The ToolContext fields the host transports to an extension child (workspace, contract) and its ABI check."""

    def __init__(self, workspace_root=None, workspace_mode="external", task_contract=None, task_metadata=None,
                 admitted=True, repo_dir=None):
        self.workspace_root = pathlib.Path(workspace_root) if workspace_root else None
        self.workspace_mode = workspace_mode
        self.task_contract = task_contract or {}
        self.task_metadata = task_metadata or {}
        self.repo_dir = pathlib.Path(repo_dir) if repo_dir else pathlib.Path("/nonexistent-ouroboros-repo")
        self.admitted = admitted

    def is_workspace_mode(self):
        return self.workspace_root is not None and bool(self.workspace_mode.strip()) and self.admitted


def call(env, tool, ctx=None, **kwargs):
    raw = env["api"].tools[tool](ctx, **kwargs)
    assert isinstance(raw, str)
    return json.loads(raw)


def wire_bytes(raw: str) -> int:
    """The skill's conservative wire measure: the handler's string JSON-encoded once more, ASCII-escaped."""
    return len(json.dumps(raw).encode("utf-8"))


# --------------------------------------------------------------------------- fixture generators

def make_docx(path):
    document = docx.Document()
    document.add_heading("Quarterly report", 1)
    document.add_paragraph("Revenue grew in every region.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Region", "Growth"
    table.cell(1, 0).text, table.cell(1, 1).text = "North", "12%"
    document.save(str(path))


def make_pptx(path):
    deck = pptx.Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Launch plan"
    slide.placeholders[1].text = "Ship in October"
    deck.save(str(path))


def make_xlsx(path):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Cities"
    sheet.append(["City", "Population"])
    sheet.append(["Oslo", 709000])
    book.create_sheet("Notes").append(["second sheet"])
    book.save(str(path))


def make_xls(path):
    book = xlwt.Workbook()
    sheet = book.add_sheet("Legacy")
    sheet.write(0, 0, "Item")
    sheet.write(0, 1, "Qty")
    sheet.write(1, 0, "Pencil")
    sheet.write(1, 1, 12)
    book.save(str(path))


def make_cfb(streams: dict) -> bytes:
    """A minimal OLE2 compound file (v3, 512-byte sectors) holding small streams in its mini stream."""
    SECTOR, MINI, END, FREE, FATSECT, NOSTREAM = 512, 64, 0xFFFFFFFE, 0xFFFFFFFF, 0xFFFFFFFD, 0xFFFFFFFF
    mini, minifat, entries = bytearray(), [], []
    for name, data in streams.items():
        assert len(data) < 4096
        start, count = len(mini) // MINI, max(1, -(-len(data) // MINI))
        mini += data + b"\0" * (count * MINI - len(data))
        minifat += list(range(start + 1, start + count)) + [END]
        entries.append((name, start, len(data)))

    def pad(blob):
        return blob + b"\0" * (-len(blob) % SECTOR)

    def entry(name, kind, child, right, start, size):
        raw = name.encode("utf-16-le") + b"\0\0"
        return (raw.ljust(64, b"\0") + struct.pack("<HBBIII", len(raw), kind, 1, NOSTREAM, right, child)
                + b"\0" * 16 + struct.pack("<I", 0) + b"\0" * 16 + struct.pack("<IQ", start, size))

    minifat_blob, mini_blob = pad(struct.pack(f"<{len(minifat)}I", *minifat)), pad(bytes(mini))
    first_dir = 1
    first_minifat = first_dir + -(-(128 * (1 + len(entries))) // SECTOR)
    first_mini = first_minifat + len(minifat_blob) // SECTOR
    directory = entry("Root Entry", 5, 1, NOSTREAM, first_mini, len(mini))
    for index, (name, start, size) in enumerate(entries, start=1):
        directory += entry(name, 2, NOSTREAM, index + 1 if index < len(entries) else NOSTREAM, start, size)
    dir_blob = pad(directory)

    def chain(first, count):
        return [first + i + 1 for i in range(count - 1)] + [END]

    fat = ([FATSECT] + chain(first_dir, len(dir_blob) // SECTOR) + chain(first_minifat, len(minifat_blob) // SECTOR)
           + chain(first_mini, len(mini_blob) // SECTOR))
    fat_blob = struct.pack(f"<{SECTOR // 4}I", *(fat + [FREE] * (SECTOR // 4 - len(fat))))
    header = (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 16 + struct.pack("<HHHHH", 0x3E, 3, 0xFFFE, 9, 6)
              + b"\0" * 6 + struct.pack("<IIIIIIIII", 0, 1, first_dir, 0, 4096, first_minifat,
                                        len(minifat_blob) // SECTOR, END, 0)
              + struct.pack("<I", 0) + b"\xff" * (4 * 108))
    assert len(header) == SECTOR
    return header + fat_blob + dir_blob + minifat_blob + mini_blob


def make_msg(path):
    """An Outlook message: MAPI property streams (subject, sender, recipient, body) in an OLE container."""
    def u16(text):
        return text.encode("utf-16-le")
    path.write_bytes(make_cfb({
        "__properties_version1.0": b"\0" * 32,
        "__substg1.0_0037001F": u16("Quarterly numbers"),
        "__substg1.0_0C1F001F": u16("alice@example.org"),
        "__substg1.0_0E04001F": u16("bob@example.org"),
        "__substg1.0_1000001F": u16("Please find the numbers attached.\r\nRegards, Alice"),
    }))


# --------------------------------------------------------------------------- reading

@pytest.mark.parametrize("name, maker, expected", [
    ("report.docx", make_docx, ["Quarterly report", "Revenue grew", "North"]),
    ("deck.pptx", make_pptx, ["Launch plan", "Ship in October"]),
    ("cities.xlsx", make_xlsx, ["Cities", "Oslo", "709000", "Notes"]),
    ("legacy.xls", make_xls, ["Legacy", "Pencil", "12"]),
    ("mail.msg", make_msg, ["alice@example.org", "bob@example.org", "Quarterly numbers", "numbers attached"]),
])
def test_binary_office_formats_become_markdown(env, name, maker, expected):
    path = env["uploads"] / name
    maker(path)
    result = call(env, "read_office_file", path=str(path))
    assert result["ok"] is True, result
    assert result["format"] == name.rsplit(".", 1)[1]
    assert result["truncated"] is False and result["returned_chars"] == result["total_chars"] == len(result["text"])
    assert result["start_char"] == 0 and result["end_char"] == result["total_chars"] and "next_start_char" not in result
    assert len(result["file_sha256"]) == 64
    for fragment in expected:
        assert fragment in result["text"]


def test_text_formats_and_home_directory_are_readable(env):
    csv_path = env["home"] / "Documents" / "orders.csv"
    csv_path.write_text("item,qty\napple,3\n", encoding="utf-8")
    html_path = env["temp"] / "page.htm"
    html_path.write_text("<html><body><h1>Title</h1><p>Para <b>bold</b></p><script>x()</script></body></html>",
                         encoding="utf-8")
    csv_result = call(env, "read_office_file", path=str(csv_path))
    assert csv_result["ok"] and csv_result["format"] == "csv" and "apple" in csv_result["text"]
    assert csv_result["encoding"] == "utf-8"
    html_result = call(env, "read_office_file", path=str(html_path))
    assert html_result["ok"] and html_result["format"] == "html"
    assert "Title" in html_result["text"] and "bold" in html_result["text"] and "x()" not in html_result["text"]
    utf16 = env["uploads"] / "wide.csv"
    utf16.write_bytes("город,число\nОсло,3\n".encode("utf-16"))  # a BOM-marked UTF-16 text file is text, not binary
    wide = call(env, "read_office_file", path=str(utf16))
    assert "Осло" in wide["text"] and wide["encoding"] == "utf-16" and "﻿" not in wide["text"]
    # An ordinary owner task (a context without a selected folder) keeps reading the home directory.
    assert call(env, "read_office_file", ctx=FakeCtx(), path=str(csv_path))["ok"] is True


@pytest.mark.parametrize("name, content", [
    ("image.csv", b"\x89PNG\r\n\x1a\n" + b"\0" * 40),
    ("paper.csv", b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\nstream"),
    ("sound.html", b"RIFF\x24\x00\x00\x00WAVEfmt "),
    ("animation.csv", b"GIF89a\x01\0\x01\0\x80\0\0" + b"\0" * 20),
    ("nulls.csv", b"a,b\n1,\0\0\0\n"),
    # Real BMP and Windows EXE headers: refused by their NUL bytes, not by their first two letters.
    ("bitmap.csv", b"BM" + struct.pack("<IHHII", 70, 0, 0, 54, 40) + struct.pack("<ii", 1, 1) + b"\x01\0\x18\0" + b"\0" * 28),
    ("program.csv", b"MZ\x90\0\x03\0\0\0\x04\0\0\0\xff\xff\0\0" + b"\0" * 44 + b"PE\0\0"),
])
def test_binary_content_behind_a_text_extension_is_parse_failed(env, name, content):
    path = env["uploads"] / name
    path.write_bytes(content)
    result = call(env, "read_office_file", path=str(path))
    assert result["ok"] is False and result["error"]["code"] == "parse_failed", result


@pytest.mark.parametrize("text", ["BMI,height_cm,weight_kg\n22.5,180,73\n", "BMW,Audi\n3,4\n",
                                  "MZ,Mozambique\nZA,South Africa\n", "RIFF_size,bytes\n1,20\n",
                                  "GIF89a,format\n1,image\n"])
def test_text_that_starts_like_a_bmp_or_exe_signature_is_read(env, text):
    path = env["home"] / "Documents" / "export.csv"
    path.write_text(text, encoding="utf-8")
    result = call(env, "read_office_file", path=str(path))
    assert result["ok"] is True and result["encoding"] == "utf-8", result
    assert text.split(",", 1)[0] in result["text"] and text.splitlines()[1].split(",")[1] in result["text"]


def test_a_docx_renamed_as_html_is_not_dumped_as_text(env):
    renamed = env["uploads"] / "word.html"
    make_docx(renamed)
    assert call(env, "read_office_file", path=str(renamed))["error"]["code"] == "parse_failed"


def test_ole_containers_are_told_apart_by_structure(env):
    sheet_as_msg = env["uploads"] / "spreadsheet.msg"
    make_xls(sheet_as_msg)  # an Excel 97-2003 workbook renamed .msg
    result = call(env, "read_office_file", path=str(sheet_as_msg))
    assert result["ok"] is False and result["error"]["code"] == "parse_failed"
    assert "Excel" in result["error"]["message"] and "rename it to .xls" in result["error"]["message"]
    mail_as_xls = env["uploads"] / "mail.xls"
    make_msg(mail_as_xls)  # an Outlook message renamed .xls
    result = call(env, "read_office_file", path=str(mail_as_xls))
    assert result["ok"] is False and result["error"]["code"] == "parse_failed" and "Outlook" in result["error"]["message"]
    other = env["uploads"] / "other.msg"
    other.write_bytes(make_cfb({"SomeStream": b"not mail"}))  # a valid OLE file that is neither
    assert call(env, "read_office_file", path=str(other))["error"]["code"] == "parse_failed"
    corrupt = env["uploads"] / "corrupt.msg"
    corrupt.write_bytes(env["plugin"].OLE_MAGIC + b"\0" * 64)
    assert call(env, "read_office_file", path=str(corrupt))["error"]["code"] == "parse_failed"


def test_a_malformed_docx_fails_instead_of_falling_back_to_a_zip_listing(env):
    broken = env["uploads"] / "broken.docx"
    with zipfile.ZipFile(broken, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", "<w:document><w:body><not closed")
    result = call(env, "read_office_file", path=str(broken))
    assert result["ok"] is False and result["error"]["code"] == "parse_failed", result
    assert "word/document.xml" not in result["error"]["message"]


def test_empty_workbooks_yield_empty_text(env):
    empty_xlsx = env["uploads"] / "empty.xlsx"
    openpyxl.Workbook().save(str(empty_xlsx))
    result = call(env, "read_office_file", path=str(empty_xlsx))
    assert result["ok"] is True and result["text"] == "" and result["total_chars"] == 0 and result["truncated"] is False
    empty_xls = env["uploads"] / "empty.xls"
    book = xlwt.Workbook()
    book.add_sheet("Blank")
    book.save(str(empty_xls))
    assert call(env, "read_office_file", path=str(empty_xls))["text"] == ""
    dash = env["uploads"] / "dash.xlsx"
    book = openpyxl.Workbook()
    book.active.append(["-"])  # table scaffolding look-alike, but a real cell value
    book.save(str(dash))
    assert "-" in call(env, "read_office_file", path=str(dash))["text"]
    no_slides = env["uploads"] / "none.pptx"
    pptx.Presentation().save(str(no_slides))
    assert call(env, "read_office_file", path=str(no_slides))["text"] == ""


def test_typed_refusals_never_raise(env, monkeypatch):
    plugin = env["plugin"]
    outside = env["outside"] / "secret.docx"
    make_docx(outside)
    refused = call(env, "read_office_file", path=str(outside))
    assert refused["ok"] is False and refused["error"]["code"] == "outside_allowed_roots"
    assert str(env["uploads"].resolve()) in refused["error"]["allowed_roots"]
    assert call(env, "read_office_file", path=str(env["uploads"] / "missing.docx"))["error"]["code"] == "not_found"
    assert call(env, "read_office_file", path="")["error"]["code"] == "not_found"
    pdf = env["uploads"] / "paper.pdf"
    pdf.write_bytes(b"%PDF-1.4")
    unsupported = call(env, "read_office_file", path=str(pdf))
    assert unsupported["error"]["code"] == "unsupported_format" and "pdf_tools" in unsupported["error"]["message"]
    legacy = env["uploads"] / "old.docx"
    legacy.write_bytes(plugin.OLE_MAGIC + b"\0" * 16)
    assert call(env, "read_office_file", path=str(legacy))["error"]["code"] == "unsupported_format"
    broken = env["uploads"] / "broken.xlsx"
    broken.write_bytes(b"PK\x03\x04 not really a workbook")
    assert call(env, "read_office_file", path=str(broken))["error"]["code"] == "parse_failed"
    make_docx(env["uploads"] / "really-a-docx.xlsx")  # a valid zip of the wrong kind is not parsed as text either
    assert call(env, "read_office_file", path=str(env["uploads"] / "really-a-docx.xlsx"))["error"]["code"] == "parse_failed"
    big = env["uploads"] / "big.csv"
    big.write_text("a,b\n1,2\n", encoding="utf-8")
    monkeypatch.setattr(plugin, "MAX_FILE_BYTES", 4)
    assert call(env, "read_office_file", path=str(big))["error"]["code"] == "too_large"
    monkeypatch.setattr(plugin, "MAX_FILE_BYTES", 50 * 1024 * 1024)
    assert call(env, "read_office_file", path=str(big), max_chars="many")["error"]["code"] == "invalid_argument"
    assert call(env, "read_office_file", path=str(big), start_char=-1)["error"]["code"] == "invalid_argument"
    assert call(env, "read_office_file", path=str(big), start_char=True)["error"]["code"] == "invalid_argument"
    monkeypatch.setattr(plugin, "_convert", lambda path, fmt, charset=None: (_ for _ in ()).throw(RuntimeError("boom")))
    failed = call(env, "read_office_file", path=str(big))
    assert failed["error"]["code"] == "parse_failed" and "boom" in failed["error"]["message"]
    monkeypatch.setattr(plugin, "read_office_file", lambda *a, **k: (_ for _ in ()).throw(ValueError("unexpected")))
    assert call(env, "read_office_file", path=str(big))["error"]["code"] == "internal_error"


def test_a_missing_converter_package_is_dependency_missing(env, monkeypatch):
    from markitdown._exceptions import MissingDependencyException

    path = env["uploads"] / "report.docx"
    make_docx(path)

    def missing(path, fmt, charset=None):
        raise MissingDependencyException("DocxConverter recognized the input but the dependencies needed to read "
                                         ".docx files have not been installed")

    monkeypatch.setattr(env["plugin"], "_convert", missing)
    result = call(env, "read_office_file", path=str(path))
    assert result["ok"] is False and result["error"]["code"] == "dependency_missing", result
    monkeypatch.setattr(env["plugin"], "_ole_kind", lambda path: (_ for _ in ()).throw(ImportError("No module named olefile")))
    make_msg(env["uploads"] / "mail.msg")
    assert call(env, "read_office_file", path=str(env["uploads"] / "mail.msg"))["error"]["code"] == "dependency_missing"


def test_start_char_reads_the_rest_of_a_long_text(env):
    path = env["uploads"] / "long.csv"
    path.write_text("text\n" + "\n".join("строка номер %d с \"кавычками\"" % i for i in range(600)), encoding="utf-8")
    full = call(env, "read_office_file", path=str(path))
    assert full["ok"] and full["truncated"] is False
    pieces, start = [], 0
    while True:
        part = call(env, "read_office_file", path=str(path), max_chars=5000, start_char=start)
        assert part["ok"] and part["file_sha256"] == full["file_sha256"] and part["total_chars"] == full["total_chars"]
        assert part["start_char"] == start and part["end_char"] == start + part["returned_chars"]
        pieces.append(part["text"])
        if not part["truncated"]:
            assert "next_start_char" not in part
            break
        assert part["returned_chars"] == 5000 and part["next_start_char"] == part["end_char"]
        start = part["next_start_char"]
    assert len(pieces) > 2 and "".join(pieces) == full["text"]
    tail = call(env, "read_office_file", path=str(path), start_char=full["total_chars"])
    assert tail["ok"] and tail["text"] == "" and tail["truncated"] is False
    beyond = call(env, "read_office_file", path=str(path), start_char=full["total_chars"] + 1)
    assert beyond["error"]["code"] == "invalid_argument" and beyond["error"]["total_chars"] == full["total_chars"]


def test_truncation_is_bounded_by_max_chars_and_host_wire_bytes(env, monkeypatch):
    plugin = env["plugin"]
    path = env["uploads"] / "long.csv"
    # Backslashes and Cyrillic are the expensive cases: each is escaped twice in the skill's conservative measure.
    path.write_text("text\n" + "\n".join("\\\\\\\\ строка %d \"кавычки\"" % i for i in range(20000)), encoding="utf-8")
    full = call(env, "read_office_file", path=str(path), max_chars=plugin.MAX_MAX_CHARS)
    assert full["ok"] and full["truncated"] is True and full["returned_chars"] < full["total_chars"]
    raw = env["api"].tools["read_office_file"](None, path=str(path), max_chars=plugin.MAX_MAX_CHARS)
    assert wire_bytes(raw) <= plugin.RESULT_BYTE_CAP < 512 * 1024
    short = call(env, "read_office_file", path=str(path), max_chars=120)
    assert short["truncated"] is True and short["returned_chars"] == 120 == len(short["text"])
    assert short["next_start_char"] == 120 and full["text"].startswith(short["text"])
    monkeypatch.setattr(plugin, "RESULT_BYTE_CAP", 3000)
    bounded = call(env, "read_office_file", path=str(path))
    raw = env["api"].tools["read_office_file"](None, path=str(path))
    assert wire_bytes(raw) <= 3000
    assert bounded["truncated"] is True and 0 < bounded["returned_chars"] < bounded["total_chars"]
    assert full["text"].startswith(bounded["text"])
    monkeypatch.setattr(plugin, "RESULT_BYTE_CAP", 10)
    assert call(env, "read_office_file", path=str(path))["text"] == ""


def make_text_docx(path, text):
    document = docx.Document()
    document.add_paragraph(text)
    document.save(str(path))


@pytest.fixture
def project(env):
    """A selected task folder on a mounted volume: outside home, temp and the data root."""
    folder = env["outside"] / "volume" / "project"
    (folder / "docs").mkdir(parents=True)
    make_text_docx(folder / "report.docx", "report from the task folder")
    make_text_docx(folder / "docs" / "notes.docx", "notes in a subfolder")
    return folder


def test_relative_paths_start_at_the_task_folder_not_the_process_cwd(env, project, monkeypatch, tmp_path):
    elsewhere = env["home"] / "cwd"
    elsewhere.mkdir()
    make_text_docx(elsewhere / "report.docx", "decoy in the process working directory")
    monkeypatch.chdir(elsewhere)  # the host starts extension children in its own repository, never the task folder
    ctx = FakeCtx(project)
    result = call(env, "read_office_file", ctx=ctx, path="report.docx")
    assert result["ok"] is True, result
    assert result["path"] == str((project / "report.docx").resolve()) and "task folder" in result["text"]
    nested = call(env, "read_office_file", ctx=ctx, path="docs/../docs/notes.docx")
    assert nested["ok"] and "subfolder" in nested["text"]
    for context in (None, FakeCtx(), FakeCtx(project, workspace_mode=""), FakeCtx(project, admitted=False)):
        refused = call(env, "read_office_file", ctx=context, path="report.docx")
        assert refused["ok"] is False and refused["error"]["code"] == "invalid_argument", refused
        assert "absolute path" in refused["error"]["message"]


def test_an_absolute_path_in_a_task_folder_outside_home_is_readable(env, project):
    target = str(project / "report.docx")
    assert call(env, "read_office_file", path=target)["error"]["code"] == "outside_allowed_roots"  # no task folder
    result = call(env, "read_office_file", ctx=FakeCtx(project), path=target)
    assert result["ok"] and "task folder" in result["text"]
    # A folder the host does not admit (no workspace mode, or one overlapping Ouroboros data) adds nothing.
    for context in (FakeCtx(project, workspace_mode=""), FakeCtx(project, admitted=False)):
        refused = call(env, "read_office_file", ctx=context, path=target)
        assert refused["error"]["code"] == "outside_allowed_roots"
        assert str(project.resolve()) not in refused["error"]["allowed_roots"]
    neighbour = env["outside"] / "volume" / "other.docx"
    make_text_docx(neighbour, "next to the task folder")
    refused = call(env, "read_office_file", ctx=FakeCtx(project), path=str(neighbour))
    assert refused["error"]["code"] == "outside_allowed_roots"
    assert str(project.resolve()) in refused["error"]["allowed_roots"]
    assert str(env["home"].resolve()) in refused["error"]["allowed_roots"]


def test_direct_room_folder_uses_host_metadata_and_keeps_protected_paths(env, project):
    ctx = FakeCtx(task_metadata={"_project_room_dir": str(project)})
    for path in ("report.docx", str(project / "report.docx")):
        result = call(env, "read_office_file", ctx=ctx, path=path)
        assert result["ok"] and "task folder" in result["text"]
    ctx.task_contract = {"resource_policy": {"protected_artifacts": [{"id": "room-file", "paths": ["report.docx"]}]}}
    assert call(env, "read_office_file", ctx=ctx, path="report.docx")["error"]["code"] == "task_policy_blocked"
    for metadata in ({"_project_room_dir": str(project), "_project_room_note": "unavailable"}, {}):
        assert call(env, "read_office_file", ctx=FakeCtx(task_metadata=metadata), path="report.docx")["ok"] is False
    other = env["home"] / "Documents"
    kept = FakeCtx(other, task_metadata={"_project_room_dir": str(project)})
    assert env["plugin"]._task_folder(kept) == other.resolve()


def test_workbook_string_typing_does_not_rescan_previous_rows(env, monkeypatch):
    from openpyxl.worksheet.worksheet import Worksheet
    def no_rescan(_self):
        raise AssertionError("max_row rescans all cells; use the known row index")
    monkeypatch.setattr(Worksheet, "max_row", property(no_rescan))
    book = env["plugin"]._build_workbook([
        {"name": "First", "columns": ["#N/A"], "rows": [["=literal"], [], ["#REF!"]], "column_widths": [], "freeze_header": True},
        {"name": "Second", "columns": [], "rows": [["#N/A"], [{"formula": "1+2"}]], "column_widths": [], "freeze_header": False},
    ])
    assert book["First"]["A4"].value == "#REF!" and book["First"]["A4"].data_type == "s"
    assert book["Second"]["A1"].data_type == "s" and book["Second"]["A2"].data_type == "f"


def test_symlinks_and_dot_dot_cannot_escape_the_task_folder(env, project):
    secret_dir = env["outside"] / "secrets"
    secret_dir.mkdir()
    make_text_docx(secret_dir / "secret.docx", "must stay unread")
    os.symlink(secret_dir / "secret.docx", project / "link.docx")
    os.symlink(secret_dir, project / "linked_dir")
    ctx = FakeCtx(project)
    for path in ("link.docx", str(project / "link.docx"), "linked_dir/secret.docx", "../../secrets/secret.docx"):
        refused = call(env, "read_office_file", ctx=ctx, path=path)
        assert refused["ok"] is False and refused["error"]["code"] == "outside_allowed_roots", (path, refused)
        assert "must stay unread" not in json.dumps(refused)
    inside = project / "inner.docx"
    os.symlink(project / "report.docx", inside)  # a link that stays inside the folder is an ordinary file
    assert call(env, "read_office_file", ctx=ctx, path="inner.docx")["ok"] is True


def test_a_narrowed_task_contract_is_not_widened(env, project):
    hidden = project / "hidden"
    hidden.mkdir()
    make_text_docx(hidden / "answers.docx", "protected answers")
    make_text_docx(project / "reference.docx", "black-box reference")
    personal = env["home"] / "Documents" / "private.docx"
    make_text_docx(personal, "home file this task must not read")
    contract = {"resource_policy": {"protected_artifacts": [
        {"id": "answer-key", "role": "black_box_reference", "paths": ["hidden"]},
        {"id": "reference", "role": "black_box_reference", "paths": "reference.docx", "allow": ["execute"]},
        {"id": "personal", "role": "black_box_reference", "paths": [str(personal)]},
    ]}}
    for ctx in (FakeCtx(project, task_contract=contract), FakeCtx(project, task_metadata={"task_contract": contract})):
        for path in ("hidden/answers.docx", str(hidden / "answers.docx"), "reference.docx", str(personal)):
            refused = call(env, "read_office_file", ctx=ctx, path=path)
            assert refused["ok"] is False and refused["error"]["code"] == "task_policy_blocked", (path, refused)
            assert not any(secret in json.dumps(refused, ensure_ascii=False) for secret in
                           ("protected answers", "black-box reference", "must not read"))
        assert call(env, "read_office_file", ctx=ctx, path="report.docx")["ok"] is True  # the rest stays readable
    # The ordinary owner task without that contract reads the same home file.
    assert call(env, "read_office_file", ctx=FakeCtx(), path=str(personal))["ok"] is True
    readable = {"resource_policy": {"protected_artifacts": [
        {"id": "a", "role": "black_box_reference", "paths": ["reference.docx"], "allow": ["read_bytes", "execute"]},
        {"id": "b", "role": "black_box_reference", "paths": ["hidden"], "deny": ["write", "delete"]},
        {"id": "c", "role": "fixture", "paths": [str(personal)]},
    ]}}
    for path in ("reference.docx", "hidden/answers.docx", str(personal)):
        assert call(env, "read_office_file", ctx=FakeCtx(project, task_contract=readable), path=path)["ok"] is True
    denied = {"resource_policy": {"protected_artifacts": [{"id": "d", "role": "fixture", "paths": ["report.docx"],
                                                           "deny": ["read_bytes"]}]}}
    refused = call(env, "read_office_file", ctx=FakeCtx(project, task_contract=denied), path="report.docx")
    assert refused["error"]["code"] == "task_policy_blocked" and "'d'" in refused["error"]["message"]


BACKEND = "/container-workspace"  # how a task run through a workspace executor sees its folder


def test_executor_spellings_of_protected_paths_name_the_host_files(env, project):
    hidden = project / "hidden"
    hidden.mkdir()
    make_text_docx(hidden / "answers.docx", "protected answers")
    make_text_docx(project / "solution.docx", "protected solution")
    contract = {"resource_policy": {"protected_artifacts": [
        {"id": "solution", "role": "black_box_reference", "paths": [f"{BACKEND}/solution.docx"]},
        {"id": "answers", "role": "black_box_reference", "paths": [f"{BACKEND}/hidden/"], "allow": ["execute"]},
    ]}}
    executors = (
        {"type": "docker_exec", "container_name": "box", "network": "none",
         "workspace_host_path": str(project), "workspace_backend_path": BACKEND},
        {"type": "docker_exec", "container_name": "box",
         "path_mappings": [{"host_path": str(env["outside"]), "backend_path": "/data"},
                           {"host_path": str(project), "backend_path": BACKEND}]},
    )
    for executor in executors:
        ctx = FakeCtx(project, task_contract=contract, task_metadata={"executor_ref": executor})
        for path in ("solution.docx", str(project / "solution.docx"), "hidden/answers.docx"):
            refused = call(env, "read_office_file", ctx=ctx, path=path)
            assert refused["ok"] is False and refused["error"]["code"] == "task_policy_blocked", (executor, path)
            assert "protected answers" not in json.dumps(refused) and "protected solution" not in json.dumps(refused)
        assert call(env, "read_office_file", ctx=ctx, path="report.docx")["ok"] is True  # unprotected files stay readable
    # Without an executor that spelling names nothing on this host, as for the host's own readers.
    assert call(env, "read_office_file", ctx=FakeCtx(project, task_contract=contract), path="solution.docx")["ok"] is True
    plugin = env["plugin"]
    mappings = plugin._executor_mappings(FakeCtx(project, task_metadata={"executor_ref": executors[1]}))
    assert [base for base, _ in mappings] == [BACKEND, "/data"]  # the longest backend path is tried first
    assert plugin._mapped_host_path(f"{BACKEND}/a/./b", mappings) == project.resolve() / "a" / "b"
    assert plugin._mapped_host_path(BACKEND, mappings) == project.resolve()
    assert plugin._mapped_host_path(f"{BACKEND}/../outside", mappings) is None  # a spelling that escapes maps nowhere
    assert plugin._mapped_host_path(f"{BACKEND}-other/x", mappings) is None


def test_hashing_is_withheld_separately_from_reading(env, project):
    contract = {"resource_policy": {"protected_artifacts": [
        {"id": "reference", "role": "black_box_reference", "paths": ["report.docx"], "allow": ["read_bytes"]},
        {"id": "notes", "role": "fixture", "paths": ["docs"], "deny": ["hash"]},
    ]}}
    ctx = FakeCtx(project, task_contract=contract)
    for path, artifact, words in (("report.docx", "reference", "task folder"), ("docs/notes.docx", "notes", "subfolder")):
        result = call(env, "read_office_file", ctx=ctx, path=path)
        assert result["ok"] is True and words in result["text"], result
        assert "file_sha256" not in result and result["file_sha256_withheld"] == artifact
    hashable = {"resource_policy": {"protected_artifacts": [
        {"id": "reference", "role": "black_box_reference", "paths": ["report.docx"], "allow": ["read_bytes", "hash"]}]}}
    result = call(env, "read_office_file", ctx=FakeCtx(project, task_contract=hashable), path="report.docx")
    assert len(result["file_sha256"]) == 64 and "file_sha256_withheld" not in result
    assert len(call(env, "read_office_file", ctx=FakeCtx(project), path="report.docx")["file_sha256"]) == 64


def test_csv_and_html_encodings_are_established_not_guessed(env):
    cyrillic = "город,население\nМосква,13000000\nСанкт-Петербург,5600000\n"
    path = env["uploads"] / "cities.csv"
    path.write_bytes(cyrillic.encode("cp1251"))
    refused = call(env, "read_office_file", path=str(path))
    assert refused["ok"] is False and refused["error"]["code"] == "unknown_encoding", refused  # no mojibake success
    assert "cp1251" in refused["error"]["message"]
    for name in ("cp1251", "windows-1251", "CP1251"):
        result = call(env, "read_office_file", path=str(path), encoding=name)
        assert result["ok"] and result["encoding"] == "cp1251" and "Москва" in result["text"], result
    assert call(env, "read_office_file", path=str(path), encoding="utf-8")["error"]["code"] == "parse_failed"
    for bad in ("klingon", "base64", "rot13", 1251, "undefined", "punycode"):
        assert call(env, "read_office_file", path=str(path), encoding=bad)["error"]["code"] == "invalid_argument", bad
    unlabelled = env["uploads"] / "domains.csv"
    unlabelled.write_bytes(b"xn--abc-")  # idna decodes "x" but refuses this label without naming a byte
    assert call(env, "read_office_file", path=str(unlabelled), encoding="idna")["error"]["code"] == "parse_failed"
    marked = env["uploads"] / "marked.csv"
    marked.write_bytes(b"\xef\xbb\xbf" + cyrillic.encode("utf-8"))  # Excel's UTF-8 export: the mark decides
    result = call(env, "read_office_file", path=str(marked), encoding="cp1251")
    assert result["encoding"] == "utf-8-sig" and "Москва" in result["text"] and "﻿" not in result["text"]
    declared = env["uploads"] / "declared.html"
    declared.write_bytes('<html><head><meta charset="windows-1251"></head><body><p>Привет мир</p></body></html>'
                         .encode("cp1251"))
    result = call(env, "read_office_file", path=str(declared))
    assert result["ok"] and result["encoding"] == "cp1251" and "Привет мир" in result["text"]
    undeclared = env["uploads"] / "undeclared.html"
    undeclared.write_bytes("<html><body><p>Привет мир</p></body></html>".encode("cp1251"))
    assert call(env, "read_office_file", path=str(undeclared))["error"]["code"] == "unknown_encoding"
    assert "Привет" in call(env, "read_office_file", path=str(undeclared), encoding="cp1251")["text"]
    for name, body, codec in (
            ("equiv.html", '<html><head><meta http-equiv="Content-Type" content="text/html; charset=windows-1251">'
                           '</head><body><p>Привет мир</p></body></html>'.encode("cp1251"), "cp1251"),
            ("wide.html", "<html><body><p>Привет мир</p></body></html>".encode("utf-16"), "utf-16")):
        (env["uploads"] / name).write_bytes(body)
        result = call(env, "read_office_file", path=str(env["uploads"] / name))
        assert result["ok"] and result["encoding"] == codec and "Привет мир" in result["text"], result
    for charset in ("no-such-codec", "undefined"):
        wrong = env["uploads"] / f"wrong-{charset}.html"
        wrong.write_bytes(f'<html><head><meta charset="{charset}"></head><body>Привет</body></html>'.encode("cp1251"))
        assert call(env, "read_office_file", path=str(wrong))["error"]["code"] == "unknown_encoding", charset
    report = env["uploads"] / "report.docx"
    make_docx(report)
    assert call(env, "read_office_file", path=str(report), encoding="cp1251")["error"]["code"] == "invalid_argument"
    assert "encoding" not in call(env, "read_office_file", path=str(report))


def test_numeric_arguments_must_be_whole_numbers(env):
    path = env["uploads"] / "short.csv"
    path.write_text("a,b\n1,2\n", encoding="utf-8")
    for kwargs in ({"max_chars": 2.5}, {"start_char": 1.5}, {"max_chars": float("inf")}, {"start_char": float("nan")},
                   {"max_chars": "3.0"}):
        result = call(env, "read_office_file", path=str(path), **kwargs)
        assert result["ok"] is False and result["error"]["code"] == "invalid_argument", (kwargs, result)
    whole = call(env, "read_office_file", path=str(path), max_chars=3.0, start_char="1")
    assert whole["ok"] and whole["returned_chars"] == 3 and whole["start_char"] == 1


# --------------------------------------------------------------------------- writing

def test_create_xlsx_file_writes_only_under_outputs(env):
    workbook = {"sheets": [
        {"name": "Sales", "columns": ["Product", "Qty", "Price"], "rows": [["Apple", 3, 1.25], ["Pear", None, True]],
         "column_widths": [20, 8], "freeze_header": True},
        {"name": "Raw", "rows": [[1, 2], ["x"]]},
    ]}
    result = call(env, "create_xlsx_file", filename="sales.xlsx", workbook=workbook)
    assert result["ok"] is True, result
    target = pathlib.Path(result["path"])
    assert target.is_absolute() and target.parent == (env["state"] / "outputs").resolve()
    assert result["sheets"] == ["Sales", "Raw"] and result["rows"] == 4 and result["size_bytes"] == target.stat().st_size
    book = openpyxl.load_workbook(target)
    sales = book["Sales"]
    assert [c.value for c in sales[1]] == ["Product", "Qty", "Price"] and sales[1][0].font.bold
    assert [c.value for c in sales[2]] == ["Apple", 3, 1.25] and sales["B3"].value is None and sales["C3"].value is True
    assert sales.freeze_panes == "A2" and sales.column_dimensions["A"].width == 20
    assert [c.value for c in book["Raw"][1]] == [1, 2]
    assert call(env, "create_xlsx_file", filename="sales.xlsx", workbook=workbook)["error"]["code"] == "already_exists"
    again = call(env, "create_xlsx_file", filename="sales.xlsx", workbook=workbook, overwrite=True)
    assert again["ok"] and again["path"] == result["path"]
    assert sorted(p.name for p in target.parent.iterdir()) == ["sales.xlsx"]
    assert list(env["temp"].iterdir()) == []  # openpyxl's worksheet temporaries never left the outputs directory


def test_only_cells_marked_as_formulas_become_formulas(env):
    workbook = {"sheets": [{"name": "S", "columns": ["=HYPERLINK(\"http://x.invalid\",\"h\")", "Qty"],
                            "rows": [["=cmd|' /C calc'!A0", 2], ["=1+1", 3],
                                     ["Total", {"formula": "SUM(B2:B3)"}], ["Double", {"formula": "=B4*2"}]]}]}
    result = call(env, "create_xlsx_file", filename="mixed.xlsx", workbook=workbook)
    assert result["ok"] and result["formula_cells"] == 2, result
    sheet = openpyxl.load_workbook(result["path"])["S"]
    assert [(c.value, c.data_type) for c in sheet[1]] == [("=HYPERLINK(\"http://x.invalid\",\"h\")", "s"), ("Qty", "s")]
    assert [(c.value, c.data_type) for c in sheet[2]] == [("=cmd|' /C calc'!A0", "s"), (2, "n")]  # copied text stays text
    assert [(c.value, c.data_type) for c in sheet[3]] == [("=1+1", "s"), (3, "n")]
    assert (sheet["B4"].value, sheet["B4"].data_type) == ("=SUM(B2:B3)", "f")
    assert (sheet["B5"].value, sheet["B5"].data_type) == ("=B4*2", "f")
    # Stored, not calculated: no cached result exists until a spreadsheet application recalculates the file.
    assert openpyxl.load_workbook(result["path"], data_only=True)["S"]["B4"].value is None
    read_back = call(env, "read_office_file", path=result["path"])
    assert read_back["ok"] and "| Total | NaN |" in read_back["text"] and "=cmd|' /C calc'!A0" in read_back["text"]
    plain = call(env, "create_xlsx_file", filename="plain.xlsx", workbook={"sheets": [{"name": "S", "rows": [["=A1"]]}]})
    assert plain["formula_cells"] == 0 and openpyxl.load_workbook(plain["path"])["S"]["A1"].data_type == "s"


def test_strings_spelled_like_excel_errors_stay_text(env):
    errors = ["#N/A", "#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#NUM!", "#NULL!"]
    workbook = {"sheets": [{"name": "S", "columns": ["#N/A", "Value"],
                            "rows": [[word, 1] for word in errors] + [["Sum", {"formula": "SUM(B2:B8)"}]]}]}
    result = call(env, "create_xlsx_file", filename="errors.xlsx", workbook=workbook)
    assert result["ok"] and result["formula_cells"] == 1, result
    sheet = openpyxl.load_workbook(result["path"])["S"]
    assert [(c.value, c.data_type) for c in sheet[1]] == [("#N/A", "s"), ("Value", "s")]
    assert [(sheet.cell(row=r, column=1).value, sheet.cell(row=r, column=1).data_type) for r in range(2, 9)] == [
        (word, "s") for word in errors]
    assert (sheet["B9"].value, sheet["B9"].data_type) == ("=SUM(B2:B8)", "f")  # the formula object stays a formula
    with zipfile.ZipFile(result["path"]) as archive:
        assert 't="e"' not in archive.read("xl/worksheets/sheet1.xml").decode("utf-8")  # no error cell is stored


@pytest.mark.parametrize("cell, fragment", [
    ({"formula": ""}, "formula text is empty"), ({"formula": "="}, "formula text is empty"),
    ({"formula": "   "}, "formula text is empty"), ({"formula": 5}, 'exactly {"formula"'),
    ({"formula": None}, 'exactly {"formula"'), ({"formula": "SUM(A1:A2)", "value": 3}, 'exactly {"formula"'),
    ({"value": "SUM(A1:A2)"}, 'exactly {"formula"'), ({}, 'exactly {"formula"'),
    ({"formula": "A1+" * 3000}, "formula longer than 8192"), ({"formula": "A1\x01"}, "control characters"),
    ([1, 2], "cells must be strings, numbers, booleans, null or"),
])
def test_malformed_formula_cells_are_refused(env, cell, fragment):
    result = call(env, "create_xlsx_file", filename="bad.xlsx", workbook={"sheets": [{"name": "S", "rows": [["ok", cell]]}]})
    assert result["ok"] is False and result["error"]["code"] == "invalid_workbook", result
    assert "rows[0][1]" in result["error"]["message"] and fragment in result["error"]["message"], result
    assert not (env["state"] / "outputs").exists() or not any((env["state"] / "outputs").iterdir())


def test_headers_stay_literal_and_the_workbook_wide_switch_is_refused(env):
    header = call(env, "create_xlsx_file", filename="h.xlsx",
                  workbook={"sheets": [{"name": "S", "columns": [{"formula": "SUM(A1)"}], "rows": []}]})
    assert header["error"]["code"] == "invalid_workbook" and "formulas belong in rows" in header["error"]["message"]
    workbook = {"sheets": [{"name": "S", "rows": [["=1+1"]]}]}
    for value in (True, False, "yes"):
        refused = call(env, "create_xlsx_file", filename="switch.xlsx", workbook=workbook, allow_formulas=value)
        assert refused["ok"] is False and refused["error"]["code"] == "invalid_argument"
        assert '{"formula"' in refused["error"]["message"]
    assert not (env["state"] / "outputs").exists() or not any((env["state"] / "outputs").iterdir())


def test_outputs_directory_that_is_a_symlink_is_refused(env):
    elsewhere = env["outside"] / "elsewhere"
    elsewhere.mkdir()
    os.symlink(elsewhere, env["state"] / "outputs")
    workbook = {"sheets": [{"name": "S", "rows": [[1]]}]}
    result = call(env, "create_xlsx_file", filename="escape.xlsx", workbook=workbook)
    assert result["ok"] is False and result["error"]["code"] == "output_outside_state_dir", result
    assert list(elsewhere.iterdir()) == []
    assert call(env, "create_xlsx_file", filename="escape.xlsx", workbook=workbook, overwrite=True)["error"]["code"] == "output_outside_state_dir"
    assert list(elsewhere.iterdir()) == []
    os.unlink(env["state"] / "outputs")
    (env["state"] / "outputs").mkdir()
    victim = env["outside"] / "victim.xlsx"
    victim.write_bytes(b"keep me")
    os.symlink(victim, env["state"] / "outputs" / "link.xlsx")  # a pre-existing link where the file would go
    assert call(env, "create_xlsx_file", filename="link.xlsx", workbook=workbook)["error"]["code"] == "already_exists"
    assert victim.read_bytes() == b"keep me"
    replaced = call(env, "create_xlsx_file", filename="link.xlsx", workbook=workbook, overwrite=True)
    assert replaced["ok"] and not (env["state"] / "outputs" / "link.xlsx").is_symlink()
    assert victim.read_bytes() == b"keep me"


@pytest.mark.skipif(os.name == "nt", reason="Windows has no directory handles; the documented path-based fallback")
# openpyxl abandons its ZipFile when the save fails; its finalizer then meets the already closed handle.
@pytest.mark.filterwarnings("ignore::pytest.PytestUnraisableExceptionWarning")
@pytest.mark.parametrize("overwrite", [False, True])
def test_outputs_replaced_by_a_link_during_a_call_cannot_redirect_the_write(env, monkeypatch, overwrite):
    plugin = env["plugin"]
    outputs, pinned, elsewhere = env["state"] / "outputs", env["state"] / "outputs-pinned", env["outside"] / "elsewhere"
    outputs.mkdir()
    elsewhere.mkdir()
    build = plugin._build_workbook

    def swap_then_build(sheets):  # after the checks and the open: outputs becomes a link to another directory
        outputs.rename(pinned)
        os.symlink(elsewhere, outputs)
        return build(sheets)

    monkeypatch.setattr(plugin, "_build_workbook", swap_then_build)
    result = call(env, "create_xlsx_file", filename="moved.xlsx", workbook={"sheets": [{"name": "S", "rows": [[1]]}]},
                  overwrite=overwrite)
    assert list(elsewhere.iterdir()) == []  # nothing was written through the link
    assert result["ok"] is False and result["error"]["code"] == "write_failed", result
    assert list(pinned.iterdir()) == []  # the claim, the temporary and the scratch directory were removed
    assert plugin._DIR_FD is True  # macOS and Linux: these operations went through the pinned directory handle


def test_two_writers_with_overwrite_false_cannot_both_succeed(env):
    plugin = env["plugin"]
    workbooks = [{"sheets": [{"name": "S", "rows": [[f"writer {n}"]]}]} for n in range(2)]
    with concurrent.futures.ThreadPoolExecutor(max_workers=2) as pool:
        outcomes = list(pool.map(lambda wb: plugin.create_xlsx_file("shared.xlsx", wb), workbooks))
    codes = sorted(("ok" if o["ok"] else o["error"]["code"]) for o in outcomes)
    assert codes == ["already_exists", "ok"], outcomes
    winner = next(index for index, o in enumerate(outcomes) if o["ok"])
    sheet = openpyxl.load_workbook(env["state"] / "outputs" / "shared.xlsx")["S"]
    assert sheet["A1"].value == f"writer {winner}"
    assert sorted(p.name for p in (env["state"] / "outputs").iterdir()) == ["shared.xlsx"]


def test_a_failed_save_leaves_nothing_behind(env, monkeypatch):
    plugin = env["plugin"]

    class Exploding:
        def save(self, handle):
            handle.write(b"partial")
            raise OSError("disk full")

    monkeypatch.setattr(plugin, "_build_workbook", lambda sheets: Exploding())
    result = call(env, "create_xlsx_file", filename="fail.xlsx", workbook={"sheets": [{"name": "S", "rows": []}]})
    assert result["ok"] is False and result["error"]["code"] == "write_failed" and "disk full" in result["error"]["message"]
    assert list((env["state"] / "outputs").iterdir()) == []  # no reservation, no temporary, no scratch directory
    assert list(env["temp"].iterdir()) == []
    (env["state"] / "outputs" / "dir.xlsx").mkdir()
    monkeypatch.undo()
    result = call(env, "create_xlsx_file", filename="dir.xlsx", workbook={"sheets": [{"name": "S", "rows": []}]}, overwrite=True)
    assert result["ok"] is False and result["error"]["code"] == "write_failed"
    assert sorted(p.name for p in (env["state"] / "outputs").iterdir()) == ["dir.xlsx"]


def test_without_a_state_directory_nothing_is_written_beside_the_code(env):
    plugin = env["plugin"]
    plugin.register(FakeAPI(None))
    result = call(env, "create_xlsx_file", filename="x.xlsx", workbook={"sheets": [{"name": "S", "rows": []}]})
    assert result["ok"] is False and result["error"]["code"] == "state_dir_unavailable"
    assert not (SKILL_DIR / "_state").exists()
    assert plugin._state_dir() is None and str(SKILL_DIR) not in json.dumps([str(r) for r in plugin._allowed_read_roots()])


@pytest.mark.parametrize("filename, code", [
    ("../escape.xlsx", "output_outside_state_dir"), ("..", "output_outside_state_dir"),
    ("/tmp/abs.xlsx", "output_outside_state_dir"), ("C:\\abs.xlsx", "output_outside_state_dir"),
    ("sub/dir.xlsx", "output_outside_state_dir"), ("nul\0.xlsx", "output_outside_state_dir"),
    ("D:escape.xlsx", "output_outside_state_dir"), ("report:stream.xlsx", "output_outside_state_dir"),
    ("", "invalid_filename"), ("report.docx", "invalid_filename"), ("x" * 130 + ".xlsx", "invalid_filename"),
    (".hidden.xlsx", "invalid_filename"), ("bad\x01.xlsx", "invalid_filename"),
])
def test_unsafe_filenames_are_refused_before_writing(env, filename, code):
    result = call(env, "create_xlsx_file", filename=filename, workbook={"sheets": [{"name": "S", "rows": []}]})
    assert result["ok"] is False and result["error"]["code"] == code
    assert not (env["state"] / "outputs").exists() or not any((env["state"] / "outputs").iterdir())


@pytest.mark.parametrize("workbook, fragment", [
    (None, "sheets"), ({"sheets": []}, "sheets"), ({"sheets": [1]}, "sheets[0] must be an object"),
    ({"sheets": [{"name": "a/b", "rows": []}]}, "name"), ({"sheets": [{"name": "", "rows": []}]}, "name"),
    ({"sheets": [{"name": "bad\x01name", "rows": []}]}, "sheets[0].name"),
    ({"sheets": [{"name": "A", "rows": []}, {"name": "a", "rows": []}]}, "duplicates"),
    ({"sheets": [{"name": "A", "rows": "nope"}]}, "rows"), ({"sheets": [{"name": "A", "rows": [1]}]}, "rows[0]"),
    ({"sheets": [{"name": "A", "rows": [[{"x": 1}]]}]}, "rows[0][0]"),
    ({"sheets": [{"name": "A", "rows": [[1, "x" * 32768]]}]}, "rows[0][1]: string longer than 32767"),
    ({"sheets": [{"name": "A", "rows": [["ok"], ["bad\x01char"]]}]}, "rows[1][0]: string contains control characters"),
    ({"sheets": [{"name": "A", "rows": [[float("nan")]]}]}, "rows[0][0]: number must be finite"),
    ({"sheets": [{"name": "A", "rows": [[float("inf")]]}]}, "rows[0][0]: number must be finite"),
    ({"sheets": [{"name": "A", "rows": [[9007199254740993]]}]}, "rows[0][0]: integer beyond"),
    ({"sheets": [{"name": "A", "columns": ["only"], "rows": [[1, 2]]}]}, "more cells than columns"),
    ({"sheets": [{"name": "A", "columns": [], "rows": []}]}, "columns"),
    ({"sheets": [{"name": "A", "columns": ["h" * 32768], "rows": []}]}, "columns[0]: string longer than 32767"),
    ({"sheets": [{"name": "A", "columns": ["ok", "bad\x01"], "rows": []}]}, "columns[1]: string contains control"),
    ({"sheets": [{"name": "A", "rows": [], "column_widths": [0]}]}, "column_widths"),
    ({"sheets": [{"name": "A", "rows": [], "freeze_header": "yes"}]}, "freeze_header"),
])
def test_invalid_workbooks_name_the_offending_field(env, workbook, fragment):
    result = call(env, "create_xlsx_file", filename="ok.xlsx", workbook=workbook)
    assert result["ok"] is False and result["error"]["code"] == "invalid_workbook"
    assert fragment in result["error"]["message"], result
    assert call(env, "create_xlsx_file", filename="ok.xlsx", workbook=workbook, overwrite="no")["error"]["code"] == "invalid_argument"


def test_boundary_values_round_trip_exactly(env):
    workbook = {"sheets": [{"name": "Edge", "columns": ["h" * 32767, "b", "c", "d", "e", "f"],
                            "rows": [["x" * 32767, 9007199254740992, -9007199254740992, 0.1, 1e300, "\t tab ok \n"]]}]}
    result = call(env, "create_xlsx_file", filename="edge.xlsx", workbook=workbook)
    assert result["ok"], result
    sheet = openpyxl.load_workbook(result["path"])["Edge"]
    assert sheet["A1"].value == "h" * 32767
    assert [c.value for c in sheet[2]] == ["x" * 32767, 9007199254740992, -9007199254740992, 0.1, 1e300, "\t tab ok \n"]
